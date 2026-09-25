import openpyxl
from django.core.management.base import BaseCommand, CommandError

from service.models import Category, Service


class Command(BaseCommand):
    help = "Importe ou synchronise les catégories et services depuis un fichier Excel."

    def add_arguments(self, parser):
        parser.add_argument(
            "fichier",
            type=str,
            help="Chemin vers le fichier .xlsx (ex: services-beobenere.xlsx)",
        )

    def handle(self, *args, **options):
        chemin = options["fichier"]

        try:
            wb = openpyxl.load_workbook(chemin, data_only=True)
        except FileNotFoundError:
            raise CommandError(f"Fichier introuvable : {chemin}")

        # ===== Import des catégories =====
        ws_cat = wb["Categories"]
        rows_cat = list(ws_cat.iter_rows(min_row=2, values_only=True))

        nb_cat_crees = 0
        nb_cat_maj = 0

        for row in rows_cat:
            if not row or not row[0]:
                continue
            name, slug = row[0], row[1]

            obj, created = Category.objects.update_or_create(
                slug=slug,
                defaults={"name": name},
            )
            if created:
                nb_cat_crees += 1
            else:
                nb_cat_maj += 1

        self.stdout.write(
            self.style.SUCCESS(
                f"Catégories : {nb_cat_crees} créée(s), {nb_cat_maj} mise(s) à jour."
            )
        )

        # ===== Import des services =====
        ws_svc = wb["Services"]
        rows_svc = list(ws_svc.iter_rows(min_row=2, values_only=True))

        nb_svc_crees = 0
        nb_svc_maj = 0
        erreurs = []

        for row in rows_svc:
            if not row or not row[0]:
                continue

            name, slug, brief_description, description, image_url, is_active_str, category_slug = row

            try:
                category = Category.objects.get(slug=category_slug)
            except Category.DoesNotExist:
                erreurs.append(
                    f"  - Service '{name}' ignoré : catégorie '{category_slug}' introuvable."
                )
                continue

            is_active = str(is_active_str).strip().upper() in ("OUI", "TRUE", "1", "YES")

            obj, created = Service.objects.update_or_create(
                slug=slug,
                defaults={
                    "name": name,
                    "brief_description": brief_description or "",
                    "description": description or "",
                    "image_url": image_url or "",
                    "is_active": is_active,
                    "category": category,
                },
            )
            if created:
                nb_svc_crees += 1
            else:
                nb_svc_maj += 1

        self.stdout.write(
            self.style.SUCCESS(
                f"Services : {nb_svc_crees} créé(s), {nb_svc_maj} mis à jour."
            )
        )

        if erreurs:
            self.stdout.write(self.style.WARNING("Avertissements :"))
            for e in erreurs:
                self.stdout.write(self.style.WARNING(e))
