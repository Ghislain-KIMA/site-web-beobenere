import openpyxl
from django.core.exceptions import ValidationError
from django.core.management.base import BaseCommand, CommandError

from company.models import Company

CHAMPS = [
    "name", "slogan", "description", "sector", "logo_url",
    "email", "phone", "address", "whatsapp", "facebook_url",
    "business_hours", "service_area",
]


class Command(BaseCommand):
    help = "Importe ou met à jour les informations de l'entreprise depuis un fichier Excel."

    def add_arguments(self, parser):
        parser.add_argument(
            "fichier",
            type=str,
            help="Chemin vers le fichier .xlsx (ex: data/services-beobenere.xlsx)",
        )

    def handle(self, *args, **options):
        chemin = options["fichier"]

        try:
            wb = openpyxl.load_workbook(chemin, data_only=True)
        except FileNotFoundError:
            raise CommandError(f"Fichier introuvable : {chemin}")

        nom_feuille = next(
            (s for s in wb.sheetnames if s.lower() == "company"), None
        )
        if nom_feuille is None:
            raise CommandError("Aucune feuille nommée 'company' dans ce fichier.")

        ws = wb[nom_feuille]
        entetes = [cell.value for cell in ws[1]]
        valeurs = [cell.value for cell in ws[2]]
        data = dict(zip(entetes, valeurs))

        company = Company.objects.first()
        created = company is None
        if created:
            company = Company()

        for champ in CHAMPS:
            setattr(company, champ, data.get(champ) or "")

        try:
            company.full_clean()
        except ValidationError as e:
            raise CommandError(
                f"Données invalides dans la feuille 'company', import annulé : {e}"
            )

        company.save()

        if created:
            self.stdout.write(self.style.SUCCESS("Entreprise créée."))
        else:
            self.stdout.write(self.style.SUCCESS("Entreprise mise à jour."))
